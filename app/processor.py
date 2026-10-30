import csv
import hashlib
import io
import math
import re
from decimal import Decimal, InvalidOperation
from typing import Any
from .config import TESSERACT_CMD

def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()

def _text_from_pdf(data: bytes) -> tuple[str, int]:
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        page_text = [page.extract_text() or "" for page in reader.pages]
        if all(text.strip() for text in page_text):
            return "\n".join(page_text), len(page_text)
        ocr_text, _ = _ocr_pdf(data, [index for index, text in enumerate(page_text) if not text.strip()])
        ocr_pages = iter(ocr_text)
        merged = [next(ocr_pages) if not text.strip() else text for text in page_text]
        return "\n".join(merged), len(page_text)
    except Exception as exc:
        raise ValueError(f"PDF extraction failed: {exc}") from exc

def _ocr_pdf(data: bytes, page_indexes: list[int] | None = None) -> tuple[list[str], int]:
    try:
        import fitz
        from PIL import Image
        import pytesseract
        if TESSERACT_CMD:
            pytesseract.pytesseract.tesseract_cmd = TESSERACT_CMD
        with fitz.open(stream=data, filetype="pdf") as pdf:
            pages = []
            targets = set(page_indexes) if page_indexes is not None else set(range(len(pdf)))
            for index, page in enumerate(pdf):
                if index not in targets:
                    continue
                pixels = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                with Image.frombytes("RGB", [pixels.width, pixels.height], pixels.samples) as image:
                    pages.append(pytesseract.image_to_string(image))
            return pages, len(pdf)
    except Exception as exc:
        raise ValueError(f"Scanned PDF OCR unavailable or failed: {exc}") from exc

def _text_from_image(data: bytes) -> tuple[str, int]:
    try:
        from PIL import Image, ImageOps, ImageSequence
        import pytesseract
        if TESSERACT_CMD:
            pytesseract.pytesseract.tesseract_cmd = TESSERACT_CMD
        pages = []
        with Image.open(io.BytesIO(data)) as source:
            for frame in ImageSequence.Iterator(source):
                with frame.convert("RGB") as image:
                    with image.resize((image.width * 2, image.height * 2)) as enlarged:
                        with ImageOps.autocontrast(enlarged) as prepared:
                            pages.append(pytesseract.image_to_string(prepared))
        return "\n".join(pages), len(pages)
    except Exception as exc:
        raise ValueError(f"Image OCR unavailable or failed: {exc}") from exc

def _text_from_sheet(data: bytes, suffix: str) -> tuple[str, int]:
    try:
        if suffix == ".csv":
            rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
            return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows), 1
        if suffix == ".xls":
            import xlrd
            workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
            chunks = []
            for sheet in workbook.sheets():
                chunks.append(f"[Sheet: {sheet.name}]")
                for row_index in range(sheet.nrows):
                    chunks.append(" | ".join(str(value) for value in sheet.row_values(row_index)))
            count = workbook.nsheets
            workbook.release_resources()
            return "\n".join(chunks), count
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        chunks = []
        for sheet in workbook.worksheets:
            chunks.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                chunks.append(" | ".join("" if value is None else str(value) for value in row))
        count = len(workbook.worksheets)
        workbook.close()
        return "\n".join(chunks), count
    except Exception as exc:
        raise ValueError(f"Spreadsheet parsing failed: {exc}") from exc

def extract_text(data: bytes, suffix: str) -> tuple[str, int]:
    if suffix == ".pdf":
        return _text_from_pdf(data)
    if suffix in {".jpg", ".jpeg", ".png", ".tif", ".tiff"}:
        return _text_from_image(data)
    if suffix in {".csv", ".xlsx", ".xls"}:
        return _text_from_sheet(data, suffix)
    raise ValueError(f"Unsupported file extension: {suffix}")

MONEY = r"(?P<value>[-+]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d{1,2})?)(?![\w.,])"
def _money(text: str, labels: list[str]) -> float | None:
    for label in labels:
        match = re.search(rf"(?<![\w-])(?:{label})(?:[ \t]*[:#.|][ \t]*|[ \t]+)\$?[ \t]*{MONEY}", text, re.I)
        if match:
            try:
                value = float(match.group("value").replace(",", ""))
                return value if math.isfinite(value) else None
            except ValueError:
                return None
    return None

def _field(text: str, labels: list[str]) -> str | None:
    for label in labels:
        match = re.search(rf"(?<!\w)(?:{label})(?:[ \t]*[:#.\-|][ \t]*|[ \t]+)([^\r\n]+)", text, re.I)
        if match:
            value = match.group(1).splitlines()[0].strip()
            return value or None
    return None

def _structured_invoice(text: str) -> dict[str, Any]:
    for line_number, line in enumerate(text.splitlines()):
        headers = [part.strip().lower().replace("_", " ") for part in line.split("|")]
        known_headers = {"invoice number", "invoice date", "supplier name", "customer name",
                         "subtotal", "tax", "total", "currency"}
        if len(known_headers.intersection(headers)) < 2:
            continue
        data_lines = text.splitlines()[line_number + 1:]
        data = next((candidate for candidate in data_lines if "|" in candidate and candidate.strip()), None)
        if data is None:
            return {}
        values = [part.strip() for part in data.split("|")]
        if len(values) < len(headers):
            return {}
        row = dict(zip(headers, values))
        return {
            "invoice_number": row.get("invoice number") or None,
            "invoice_date": row.get("invoice date") or None,
            "supplier_name": row.get("supplier name") or None,
            "customer_name": row.get("customer name") or None,
            "subtotal": _parse_money_value(row.get("subtotal")),
            "tax": _parse_money_value(row.get("tax")),
            "total": _parse_money_value(row.get("total")),
            "currency": row.get("currency") or None,
        }
    return {}

def _parse_money_value(value: str | None) -> float | None:
    if not value:
        return None
    value = value.strip().removeprefix("$").strip()
    if not re.fullmatch(MONEY, value):
        return None
    try:
        parsed = float(value.replace(",", "").replace("$", "").strip())
        return parsed if math.isfinite(parsed) else None
    except ValueError:
        return None

def extract_invoice(text: str) -> tuple[dict[str, Any], float]:
    invoice: dict[str, Any] = {
        "invoice_number": _field(text, [r"invoice\s*(?:number|no\.?|#)"]),
        "invoice_date": _field(text, [r"invoice\s*date", r"date"]),
        "supplier_name": _field(text, [r"supplier", r"vendor", r"from"]),
        "customer_name": _field(text, [r"customer", r"bill\s*to", r"to"]),
        "subtotal": _money(text, [r"subtotal", r"sub-total"]),
        "tax": _money(text, [r"tax", r"vat"]),
        "total": _money(text, [r"grand\s*total", r"total\s*amount", r"total"]),
        "currency": _field(text, [r"currency"]),
        "line_items": [],
    }
    invoice.update(_structured_invoice(text))
    for line in text.splitlines():
        match = re.match(r"\s*(\d+)\s*[|,\t]\s*(.+?)\s*[|\],\t]\s*(\d+(?:\.\d+)?)\s*[|\],\t]\s*\$?([\d,]+(?:\.\d{1,2})?)\s*$", line)
        if match:
            invoice["line_items"].append({"quantity": float(match.group(1)), "description": match.group(2).strip(),
                                          "unit_price": float(match.group(3)), "amount": float(match.group(4).replace(",", ""))})
    values = sum(value is not None for key, value in invoice.items() if key != "line_items")
    return invoice, round(values / 8, 2)

def _decimal(value: Any) -> Decimal | None:
    try:
        parsed = Decimal(str(value))
        return parsed if parsed.is_finite() else None
    except (InvalidOperation, TypeError, ValueError):
        return None

def validate_invoice(invoice: dict[str, Any], text: str, document_hash: str | None = None) -> list[dict[str, str]]:
    errors = []
    required = ["invoice_number", "invoice_date", "supplier_name", "customer_name", "subtotal", "tax", "total", "currency"]
    for field in required:
        if invoice.get(field) in (None, ""):
            errors.append({"code": "missing_field", "message": f"Missing required field: {field}", "severity": "warning"})
    subtotal, tax, total = (_decimal(invoice.get(x)) for x in ("subtotal", "tax", "total"))
    for field in ("subtotal", "tax", "total"):
        value = invoice.get(field)
        if value not in (None, "") and _decimal(value) is None:
            errors.append({"code": "invalid_numeric", "message": f"{field} must be numeric", "severity": "error"})
    for field, value in (("subtotal", subtotal), ("tax", tax), ("total", total)):
        if value is not None and value < 0:
            errors.append({"code": "negative_numeric", "message": f"{field} cannot be negative", "severity": "error"})
    if subtotal is not None and tax is not None and total is not None and abs(subtotal + tax - total) > Decimal("0.02"):
        errors.append({"code": "total_mismatch", "message": "Subtotal plus tax does not equal total", "severity": "error"})
    line_total = Decimal(0)
    line_items = invoice.get("line_items", [])
    if not isinstance(line_items, list):
        errors.append({"code": "invalid_line_item", "message": "line_items must be an array", "severity": "error"})
        line_items = []
    for index, item in enumerate(line_items):
        if not isinstance(item, dict):
            errors.append({"code": "invalid_line_item", "message": f"Line item {index + 1} must be an object", "severity": "error"})
            continue
        amount = _decimal(item.get("amount"))
        quantity = _decimal(item.get("quantity"))
        unit_price = _decimal(item.get("unit_price"))
        if quantity is None or unit_price is None:
            errors.append({"code": "invalid_line_item", "message": f"Line item {index + 1} quantity and unit price must be numeric", "severity": "error"})
        if amount is None:
            errors.append({"code": "invalid_line_item", "message": f"Line item {index + 1} amount must be numeric", "severity": "error"})
            continue
        if amount < 0 or (quantity is not None and quantity < 0) or (unit_price is not None and unit_price < 0):
            errors.append({"code": "negative_line_item", "message": f"Line item {index + 1} cannot be negative", "severity": "error"})
        line_total += amount
        if quantity is not None and unit_price is not None and abs(quantity * unit_price - amount) > Decimal("0.02"):
            errors.append({"code": "line_item_math_mismatch", "message": f"Line item {index + 1} quantity times unit price does not equal amount", "severity": "error"})
    if line_items and subtotal is not None and abs(line_total - subtotal) > Decimal("0.02"):
        errors.append({"code": "line_items_mismatch", "message": "Line item amounts do not equal subtotal", "severity": "error"})
    return errors

def compare_invoices(left: dict[str, Any], right: dict[str, Any]) -> list[dict[str, str]]:
    mismatches = []
    for field in ("invoice_number", "invoice_date", "supplier_name", "customer_name",
                  "subtotal", "tax", "total", "currency"):
        if left.get(field) != right.get(field):
            mismatches.append({"field": field, "left": str(left.get(field)), "right": str(right.get(field))})
    if left.get("line_items") != right.get("line_items"):
        mismatches.append({"field": "line_items", "left": "different", "right": "different"})
    return mismatches
